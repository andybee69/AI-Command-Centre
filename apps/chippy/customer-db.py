import argparse, json, os, sys
from openpyxl import load_workbook

DB = r"C:\Users\andyb\OneDrive - Agon\Agon Info\Agon_Master folder\Database\2026\MASTER_Customers_2026_v2.xlsx"
SHEET = "ALL CUSTOMERS 2026"

def clean(v):
    return "" if v is None else str(v).strip()

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--area",default="")
    ap.add_argument("--type",default="")
    ap.add_argument("--q",default="")
    ap.add_argument("--limit",type=int,default=200)
    args=ap.parse_args()
    wb=load_workbook(DB,read_only=True,data_only=True)
    ws=wb[SHEET]
    rows=ws.iter_rows(values_only=True)
    headers=next(rows)
    idx={h:i for i,h in enumerate(headers)}
    area=args.area.casefold().strip()
    typ=args.type.casefold().strip()
    q=args.q.casefold().strip()
    out=[]; seen=set(); types=set()
    for r in rows:
        cust_type=clean(r[idx["Customer Type"]])
        if cust_type: types.add(cust_type)
        company=clean(r[idx["Company"]])
        city=clean(r[idx["City"]])
        county=clean(r[idx["County"]])
        postcode=clean(r[idx["Postcode"]]).upper()
        country=clean(r[idx["Country"]])
        street=clean(r[idx["Street"]])
        website=clean(r[idx["Website"]])
        first=clean(r[idx["First Name"]]); last=clean(r[idx["Last Name"]])
        email=clean(r[idx["Email"]]); phone=clean(r[idx["Phone"]]); mobile=clean(r[idx["Mobile"]])
        hay_area=" ".join([city,county,postcode,country]).casefold()
        hay_q=" ".join([company,first,last,email,website,street,city,postcode,cust_type]).casefold()
        if area and area not in hay_area: continue
        if typ and typ != cust_type.casefold(): continue
        if q and q not in hay_q: continue
        key=(company.casefold(),street.casefold(),city.casefold(),postcode)
        if key in seen: continue
        seen.add(key)
        out.append({
            "company":company,"firstName":first,"lastName":last,"email":email,
            "phone":phone or mobile,"street":street,"city":city,"county":county,
            "postcode":postcode,"country":country,"website":website,"customerType":cust_type
        })
        if len(out)>=max(1,min(args.limit,500)): break
    print(json.dumps({"source":DB,"count":len(out),"types":sorted(types,key=str.casefold),"items":out},ensure_ascii=False))

if __name__=="__main__":
    main()
